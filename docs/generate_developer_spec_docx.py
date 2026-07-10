#!/usr/bin/env python3
"""Generate VideoDating developer specification DOCX (companion to Miro board)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUTPUT = Path(__file__).resolve().parent / "VideoDating-Developer-Spec.docx"
MIRO_URL = "https://miro.com/app/board/uXjVH9Gr3u0=/"


def set_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph()


def build(doc: Document) -> None:
    title = doc.add_heading("VideoDating — спецификация для разработчиков", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    para(doc, "Этап 0: проектирование. Документ дополняет интерактивную карту Miro.")
    para(doc, f"Карта Miro: {MIRO_URL}")
    para(doc, "Платформа: мобильное приложение iOS / Android. Сайт — рекламный лендинг, не веб-версия app.")

    heading(doc, "1. User Flow и логика анкет")
    para(doc, "Воронка:")
    bullets(doc, [
        "Первый визит → гостевой режим (пол + возраст) → лента",
        "Регистрация → видео-визитка → модерация",
        "Одобрение → полный доступ в ленте → лайки → мэтч → чат → видео-свидание → оценка",
    ])

    heading(doc, "1.1. Гостевой режим", 2)
    table(doc, ["Действие", "Поведение"], [
        ["Первый запуск", "Пол + возраст → сразу лента"],
        ["Лайки", "До 5 за сессию, копятся, получатель не видит"],
        ["Корзина «Лайкнули меня»", "Заблокирована"],
        ["Чаты", "Недоступны"],
    ])
    para(doc, "Edge case: 5 лайков исчерпаны → экран «Зарегистрируйтесь» + показ накопленных симпатий.")

    heading(doc, "1.2. Регистрация и модерация", 2)
    bullets(doc, [
        "Обязательно: имя, дата рождения, город, email/телефон + пароль, видео-визитка ≤60 сек, согласие 152-ФЗ",
        "Визитка: вопросы каждые 15–20 сек, авто стоп-кадр, статус «На проверке»",
        "На модерации: можно смотреть чужие анкеты и лайкать (30/сутки); нельзя — своя карточка в ленте и чаты",
        "Отклонение: причина + «Перезаписать визитку»",
    ])

    heading(doc, "1.3. Структура анкеты (глубокая)", 2)
    para(doc, "Поля для фильтрации ленты и совместимости:")
    table(doc, ["Поле", "Тип", "Индексация", "Назначение"], [
        ["gender", "enum", "btree", "Базовый подбор"],
        ["birth_date", "date", "возраст", "Фильтр по возрасту"],
        ["city", "string + geo_id", "btree/geo", "Город и радиус"],
        ["education", "enum/string", "btree", "Образование"],
        ["goals", "enum[]", "GIN", "Цели знакомства"],
        ["interests", "enum[]", "GIN", "Интересы (справочник)"],
        ["has_children", "bool", "btree", "С детьми / без"],
        ["moderation_status", "enum", "btree", "approved / pending / rejected"],
        ["profile_frozen", "bool", "partial", "Скрытые профили"],
    ])
    para(doc, "Справочники целей, интересов, образования — фиксированные ID, не свободный текст.")

    heading(doc, "1.4. Навигация", 2)
    bullets(doc, ["Анкеты", "Видео-лента", "Сообщения", "Корзина", "Мой профиль"])

    heading(doc, "2. Лента рекомендаций")
    bullets(doc, [
        "Фильтрация по совместимости: пол, возраст, город, цели, интересы, образование",
        "Два канала: Анкеты (свайп) и Видео-лента (рилсы) без дублей между каналами",
        "Сортировка с учётом рейтинга и приоритетов",
    ])

    heading(doc, "3. Лайки, корзина, мэтчи")
    table(doc, ["Параметр", "Значение"], [
        ["Лимит гостя", "5 лайков"],
        ["Лимит зарегистрированного", "30 / сутки"],
        ["Взаимность", "Автооткрытие чата"],
        ["Контакты", "Только после взаимного лайка"],
    ])

    heading(doc, "4. Чаты")
    table(doc, ["Правило", "Значение"], [
        ["Доступ", "Регистрация + мэтч + одобренная визитка"],
        ["Макс. активных чатов", "5"],
        ["Формат", "Только видео-кружки ≤60 сек, без текста"],
        ["Лимит от одного", "10 кружков"],
        ["Архивация", "Лимит кружков или 7 дней неактивности"],
    ])
    bullets(doc, ["Записать кружок", "Предложить видео-свидание", "Профиль собеседника"])

    heading(doc, "5. Видео-свидания (WebRTC)")
    bullets(doc, [
        "«Прямо сейчас» — подтверждение за 5 минут",
        "«По расписанию» — пересечение слотов",
        "Во время звонка: переключение камер, таймер, поминутное списание",
        "Оплата: мужчина 75%, женщина 25%; опция «оплачу полностью»; при равном выборе — у кого больше баланс",
        "Запись: согласие обоих, хранение 30 дней (платно)",
    ])
    para(doc, "Технически: WebSocket-сигналинг, WebRTC P2P, TURN при необходимости.")

    heading(doc, "6. Умная модерация")
    para(doc, "Pipeline:")
    bullets(doc, [
        "1. Загрузка видео",
        "2. FFmpeg: формат, ≤60 сек, размер",
        "3. AI пре-модерация: человек в кадре, NSFW, score 0–100",
        "4. Решение по score",
        "5. Ручная модерация (админ)",
        "6. Публикация в ленте",
    ])
    table(doc, ["Score", "Действие"], [
        ["≥ 85", "Очередь админа, «вероятно ОК»"],
        ["50–84", "Обязательный ручной просмотр"],
        ["< 50", "Автоотклонение + причина"],
    ])
    para(doc, "~70–80% отсекается автоматически. SLA для пользователя: до 24 ч.")

    heading(doc, "7. Личный кабинет, статистика, рейтинги")
    heading(doc, "7.1. ЛК", 2)
    table(doc, ["Блок", "Содержание"], [
        ["Публичная карточка", "Фото, имя, возраст, город, привлекательность"],
        ["Рейтинг", "4 критерия + общий балл"],
        ["Совместимость %", "Дата/время/место рождения"],
        ["Статистика", "Просмотры анкеты, лайки, мэтчи, свидания, отзывы"],
        ["Динамика рейтинга", "Средние по 4 критериям, влияние на ленту"],
        ["Баланс", "Пополнение, история"],
        ["Подписка", "План, окончание, смена"],
        ["Настройки", "Приватность, уведомления, заморозка/удаление"],
    ])

    heading(doc, "7.2. Оценки после свидания", 2)
    bullets(doc, [
        "4 критерия (1–5): Адекватность, Чувство юмора, Доброта, Эмпатия",
        "Текстовый отзыв (опционально), средние видны всем",
        "1 отзыв на свидание (unique video_date_id + rater_id)",
        "Только после завершённого звонка",
        "Аномалии (10× «5» за час) → флаг админу",
    ])

    heading(doc, "8. Монетизация")
    bullets(doc, [
        "Бесплатно: гость, регистрация, первые лайки",
        "1-й платёж: баланс перед свиданием",
        "Звонок: поминутно (пакеты + поминутный тариф в pricing_plans)",
        "Запись 30 дней — платно",
        "Платежи через шлюз (PCI), пополнение с 2FA",
    ])

    heading(doc, "9. Рекламный сайт (лендинг)")
    bullets(doc, [
        "Hero: УТП + видео + CTA скачать",
        "Как работает, преимущества, отзывы, тарифы, FAQ",
        "Footer: App Store / Google Play, ПДн, оферта",
        "Сайт ≠ приложение — только призыв к скачиванию",
    ])

    heading(doc, "10. Схема БД (ключевые сущности)")
    bullets(doc, [
        "users, user_sessions, likes, chats, chat_messages",
        "video_profiles, moderation_queue",
        "ratings, video_dates, pricing_plans, transactions",
        "user_goals, user_interests, education_levels (справочники)",
    ])

    heading(doc, "11. API-контуры")
    table(doc, ["Модуль", "Эндпоинты (контур)"], [
        ["Auth", "register, login, refresh, guest session, 2FA"],
        ["Profile", "me, upload-url, confirm video, questions"],
        ["Feed", "cards, reels"],
        ["Basket", "likes sent/received, mutual"],
        ["Chats", "list, circles, archive"],
        ["Calls", "initiate, schedule, billing webhook"],
        ["Ratings", "submit after call"],
        ["Balance", "top-up, history, packages"],
    ])

    heading(doc, "12. Безопасность")
    bullets(doc, [
        "Сессии: refresh SHA-256 в user_sessions, ротация",
        "JWT access + session_id, Argon2id пароли",
        "Rate limit: login 5/мин, register 3/час",
        "152-ФЗ: consent timestamps, audit log, версия согласия",
        "2FA TOTP для платежей",
    ])

    heading(doc, "13. Этапы реализации")
    table(doc, ["#", "Содержание", "Статус"], [
        ["1", "Auth, гость, feed", "Спроектировано"],
        ["2", "Анкеты, видео-лента", "Спроектировано"],
        ["3", "Корзина, лайки, мэтчи", "Спроектировано"],
        ["4", "Чаты, кружки", "Спроектировано"],
        ["5", "WebRTC, биллинг", "Спроектировано"],
        ["6", "Оценки", "Спроектировано"],
        ["7", "Профиль, ЛК, статистика", "Спроектировано"],
        ["8", "AI-модерация, админка", "Спроектировано"],
        ["9", "Лендинг", "Отдельный трек"],
        ["10", "Релиз в сторы", "Планируется"],
    ])

    heading(doc, "14. Поставка заказчику")
    table(doc, ["Артефакт", "Описание"], [
        ["Miro-карта", "Экраны, развилки, переходы, пограничные сценарии"],
        ["Этот DOCX", "Структурированное ТЗ для backend и frontend"],
        ["Схема БД", "PostgreSQL, индексы под фильтрацию"],
        ["API", "FastAPI async, Redis, S3, Arq, WebRTC"],
    ])


def main() -> None:
    doc = Document()
    set_defaults(doc)
    build(doc)
    doc.save(OUTPUT)
    print(f"Written {OUTPUT}")


if __name__ == "__main__":
    main()
